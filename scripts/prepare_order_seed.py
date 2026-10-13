"""订单表初始数据脱敏：删除敏感列，输出与原表结构一致的 xlsx，可直接用于订单表导入。

用法（在仓库根目录执行）：
    python3 scripts/prepare_order_seed.py <原始订单表.xlsx> <输出.xlsx>

规则：
- 删除敏感列（与 backend/app/importers/order.py 的 EXCLUDED_FIELDS 保持一致）：乘机人姓名、票号、PNR、航司大编码。
- 保留其余列、sheet 名与列顺序，行值原样保留，导入器的必需列不受影响。
- 扫描剩余单元格：疑似票号（3 位航司代码 + 10 位数字）只统计数量，不打印内容；命中数大于 0 时退出码为 2，需人工复核。
- 只打印统计信息（行数、删除的列名、扫描命中数），不打印任何单元格值。
"""

import re
import sys
from pathlib import Path

import openpyxl

# 与 backend/app/importers/order.py 中的 EXCLUDED_FIELDS 保持一致，修改时两边都要改
EXCLUDED_FIELDS = {"乘机人姓名", "票号", "PNR", "航司大编码"}

# 疑似票号：3 位数字航司代码 + 10 位数字（如 781 + 10 位），前后不能紧跟数字
TICKET_RE = re.compile(r"(?<!\d)\d{3}-?\d{10}(?!\d)")
# 这些列是业务订单号，导入时用于关联，格式会与票号规则重合，扫描时豁免（票号列本身已整列删除）
SCAN_EXEMPT = {"供应商订单号"}


def sanitize_sheet(ws_in, ws_out) -> dict:
    headers = [c.value for c in next(ws_in.iter_rows(min_row=1, max_row=1))]
    keep_idx = [i for i, h in enumerate(headers) if h not in EXCLUDED_FIELDS]
    dropped = [h for h in headers if h in EXCLUDED_FIELDS]

    ws_out.append([headers[i] for i in keep_idx])
    rows = 0
    ticket_hits = 0
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        kept = [row[i] for i in keep_idx]
        for h, v in zip([headers[i] for i in keep_idx], kept):
            if h not in SCAN_EXEMPT and isinstance(v, str) and TICKET_RE.search(v):
                ticket_hits += 1
        ws_out.append(kept)
        rows += 1
    return {"行数": rows, "删除的列": dropped, "保留列数": len(keep_idx), "疑似票号命中单元格数": ticket_hits}


def main(src: str, dst: str) -> int:
    src_path, dst_path = Path(src), Path(dst)
    if not src_path.is_file():
        print(f"找不到源文件：{src_path}", file=sys.stderr)
        return 1
    dst_path.parent.mkdir(parents=True, exist_ok=True)

    # 不用 read_only：该文件的 sheet 维度信息不完整，read_only 模式会读不到数据
    wb_in = openpyxl.load_workbook(src_path, data_only=True)
    wb_out = openpyxl.Workbook(write_only=True)
    report = {}
    for ws_in in wb_in.worksheets:
        ws_out = wb_out.create_sheet(ws_in.title)
        report[ws_in.title] = sanitize_sheet(ws_in, ws_out)
    wb_out.save(dst_path)

    total_hits = 0
    for name, r in report.items():
        print(f"[{name}] 行数={r['行数']} 保留列={r['保留列数']} 删除的列={r['删除的列']} 疑似票号命中={r['疑似票号命中单元格数']}")
        total_hits += r["疑似票号命中单元格数"]
    print(f"输出：{dst_path}")
    if total_hits:
        print("存在疑似票号残留，请人工复核后再使用。", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    sys.exit(main(sys.argv[1], sys.argv[2]))
